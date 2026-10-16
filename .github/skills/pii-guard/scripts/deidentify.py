#!/usr/bin/env python3
"""
deidentify.py: make a de-identified copy of conversion extracts (CSV / XLSX).

The copy keeps structure, formats and joins intact, so reshape, mapping, rules,
demos and partner work run without real personal data.

  1) Draft a policy from column names and a Presidio sample scan (prints NO values):
       python deidentify.py suggest --in <folder or files> --out policy.draft.yaml

  2) Review the draft, then apply it:
       set PII_SALT=<secret from Key Vault>   # the same salt keeps tokens consistent across files and takes
       python deidentify.py apply --policy policy.yaml --in <folder> --out <output folder>

Guarantees:
  - Every row and column position is kept, so report layouts survive. Values are read
    and written as text, so leading zeros survive. XLSX formatting is not kept.
  - The same input value in the same domain always maps to the same output. That keeps
    account numbers, meter IDs and names joinable across files.
  - The run report (JSON) contains counts only, never values.
  - Columns the policy doesn't mention stop the run, unless the policy sets
    defaults.unknown_columns explicitly.

Needs the project virtual environment: pandas, openpyxl, pyyaml, faker, presidio-analyzer and en_core_web_lg.
"""
from __future__ import annotations

import argparse
import fnmatch
import hashlib
import hmac
import json
import os
import random
import re
import string
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
import yaml
from faker import Faker

SKIP_DEFAULT = [r"^\s*$", r"^[-=_*.\s]+$"]
LONG_NUMBER_RE = re.compile(r"\d{7,}")
DATE_FORMATS = ["%m/%d/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y %I:%M:%S %p", "%Y%m%d", "%m%d%Y",
                "%Y-%m", "%m/%Y"]
PRESIDIO_ENTITIES = ["PERSON", "PHONE_NUMBER", "EMAIL_ADDRESS", "US_SSN", "US_DRIVER_LICENSE", "LOCATION",
                     "CREDIT_CARD", "US_BANK_NUMBER", "IP_ADDRESS", "DATE_TIME"]
VOWELS, CONSONANTS = "AEIOU", "BCDFGHJKLMNPQRSTVWXYZ"
# Words kept as-is by pseudo_name / pseudo_address so that parsing rules (titles, suffixes,
# business names, street types, directions, units) behave exactly as on the real data.
NAME_KEEP = {
    "MR", "MRS", "MS", "MISS", "DR", "REV", "JR", "SR", "II", "III", "IV", "ESTATE", "EST", "OF", "THE", "AND",
    "ATTN", "CO", "C", "O", "LLC", "INC", "CORP", "CORPORATION", "LTD", "LP", "LLP", "PC", "PLLC", "TRUST",
    "TRUSTEE", "CHURCH", "CITY", "COUNTY", "SCHOOL", "DISTRICT", "ISD", "HOA", "ASSOC", "ASSOCIATION",
    "APARTMENTS", "APTS", "MANAGEMENT", "MGMT", "PROPERTIES", "PROPERTY", "HOMES", "BUILDERS", "COMPANY",
    "BANK", "FUND", "PARTNERS", "GROUP", "SERVICES", "UTILITY", "WATER", "DEPARTMENT", "DEPT", "VACANT",
    # operational markers that conversion rules test for (e.g. reads WHERE Account_Full_Name NOT LIKE '%IDLE%')
    "IDLE", "METER", "PULLED", "UNKNOWN", "OCCUPANT", "RESIDENT", "TENANT", "OWNER", "INACTIVE", "FINAL",
    "FINALED", "IRRIGATION", "IRR", "SPRINKLER", "HYDRANT", "CONSTRUCTION", "TEMP", "TEMPORARY",
}
ADDRESS_KEEP = NAME_KEEP | {
    "ST", "STREET", "AVE", "AV", "AVENUE", "DR", "DRIVE", "RD", "ROAD", "LN", "LANE", "CT", "COURT", "BLVD",
    "BOULEVARD", "CIR", "CIRCLE", "WAY", "PL", "PLACE", "TRL", "TRAIL", "PKWY", "PARKWAY", "HWY", "HIGHWAY",
    "LOOP", "TER", "TERRACE", "PT", "POINT", "SQ", "RUN", "PATH", "XING", "CV", "COVE", "BND", "BEND", "CRK",
    "HL", "HLS", "PASS", "ROW", "WALK", "PLZ", "PLAZA", "FWY", "EXPY", "N", "S", "E", "W", "NE", "NW", "SE",
    "SW", "NORTH", "SOUTH", "EAST", "WEST", "APT", "UNIT", "STE", "SUITE", "BLDG", "LOT", "SPC", "SPACE", "RM",
    "FL", "FLOOR", "PO", "P", "BOX", "RR", "HC", "FM", "CR", "US", "SH", "IH", "TX", "LOOP", "BYP",
}

# (regex on the column name, suggested spec, why)
NAME_RULES = [
    (r"ssn|social|\bsin\b|tax.?id|\bsoc\b", {"action": "blank"}, "government ID"),
    (r"birth|\bdob\b", {"action": "shift_date", "mode": "per_value", "max_days": 365, "on_unparseable": "blank"}, "date of birth"),
    (r"(\bdl\b|licen|driver).*(state|\bst\b|prov)", {"action": "keep"}, "licence state or province (the 2-letter rule needs it)"),
    (r"licen|\bdl\b|driver", {"action": "token_alnum", "domain": "licence"}, "driver's licence"),
    (r"e-?mail", {"action": "fake_email"}, "email"),
    (r"phone|\btel\b|cell|mobile|\bfax\b|home\s*(number|no\b|#)", {"action": "fake_phone"}, "phone"),
    (r"class|status|type\b|_type|code\b|_code|cycle|route|zone|timezone|unit", {"action": "keep"}, "code column"),
    (r"bank|routing|\baba\b|iban|card", {"action": "token_digits", "domain": "bank"}, "bank or card number"),
    (r"^acct|account[\s_]*(#|no|num|number|id)?$|accountnum|acct[\s_]*no|incode\s*acct", {"action": "token_digits", "preserve_prefix": 2, "domain": "account"}, "account number (join key)"),
    (r"meter.?id|meter\s*#|^meter\d?$", {"action": "token_alnum", "domain": "meter"}, "meter identifier (join key)"),
    (r"serial|_sn$|\bsn\b|endpoint.?sn", {"action": "token_alnum", "domain": "device_serial"}, "device serial number"),
    (r"location.?id|service.?point.?id|premise", {"action": "token_alnum", "domain": "location"}, "location identifier (join key)"),
    (r"\blat|\blon|latitude|longitude|gps|coord", {"action": "round", "decimals": 2}, "coordinates"),
    (r"attn|c/o|care.?of", {"action": "pseudo_name", "domain": "person"}, "care-of name"),
    (r"addr|street", {"action": "pseudo_address", "domain": "street"}, "street address"),
    (r"name|customer|owner|tenant|contact", {"action": "pseudo_name", "domain": "person"}, "person or company name"),
    (r"note|comment|remark|memo", {"action": "redact_text"}, "free text"),
]
STRONG_ENTITIES = {"PERSON", "EMAIL_ADDRESS", "PHONE_NUMBER", "US_SSN", "CREDIT_CARD", "US_BANK_NUMBER",
                   "US_DRIVER_LICENSE", "IP_ADDRESS"}


# ---------------------------------------------------------------- helpers
def load_salt():
    s = os.environ.get("PII_SALT")
    if not s:
        sys.exit("Set PII_SALT (a secret, ideally from Key Vault). Keep it the same across files and takes "
                 "so tokens stay consistent; never commit it.")
    return s.encode()


class Pseudonymizer:
    def __init__(self, salt: bytes):
        self.salt = salt
        self.cache = {}
        self.fake = Faker("en_US")
        self._analyzer = None
        self.fwd = {}          # domain -> {original key: token}
        self.rev = {}          # domain -> {token: original key}
        self.collision_fallbacks = {}

    def digest(self, domain, value, n=32):
        out, counter = b"", 0
        while len(out) < n:
            out += hmac.new(self.salt, f"{domain}\x1f{counter}\x1f{value}".encode("utf-8"), hashlib.sha256).digest()
            counter += 1
        return out[:n]

    def _injective(self, domain, key, make):
        """Return a token for key that no other key in this domain has (re-draw on collision).

        make(attempt) builds a candidate; attempt 0 is the plain keyed digest, so tokens are stable
        across runs unless a collision forced a re-draw. Equality between originals is preserved
        exactly: same original -> same token, different originals -> different tokens.
        """
        fwd = self.fwd.setdefault(domain, {})
        if key in fwd:
            return fwd[key]
        rev = self.rev.setdefault(domain, {})
        cand = None
        for attempt in range(5000):
            cand = make(attempt)
            # never map a value to itself (a short surname or house number could otherwise survive by chance)
            if cand != key and rev.get(cand, key) == key:
                break
        else:
            self.collision_fallbacks[domain] = self.collision_fallbacks.get(domain, 0) + 1
        fwd[key], rev[cand] = cand, key
        return cand

    def _keyed(self, domain, key, attempt, n=64):
        return self.digest(domain, key if attempt == 0 else f"{key}\x1e{attempt}", n)

    def rng(self, domain, value):
        return random.Random(int.from_bytes(self.digest(domain, value, 8), "big"))

    def faker_for(self, domain, value):
        self.fake.seed_instance(int.from_bytes(self.digest(domain, value, 8), "big"))
        return self.fake

    @staticmethod
    def match_case(src, out):
        letters = [c for c in src if c.isalpha()]
        if letters and all(c.isupper() for c in letters):
            return out.upper()
        return out

    def token_digits(self, v, domain, preserve_prefix=0, group_digits=0, keep_suffix=False,
                     strip_leading_zeros=False, nonzero_lead=False):
        """Replace digits with keyed digits; letters and separators stay where they are.

        Keyed on the digits only: "09-0140" and "090140" get the same token digits.
        group_digits=N   the token of the first N digits depends only on those N digits, so values that share
                         a prefix keep sharing it (account 04-0123-01: LEFT(acct, 7) = route + location group).
        keep_suffix      leave the digits after the group unchanged (an occupant sequence such as -01).
        strip_leading_zeros  key on the digits after any leading zeros and keep the zeros, so "012345678" and
                         "12345678" (Excel dropped the zero) share one core token, as they share digits.
        nonzero_lead     the first tokenized digit is never 0, so "starts with 0" stays exactly as in the source.
        """
        digits = re.sub(r"\D", "", v)
        if not digits:
            return v
        core = digits.lstrip("0") if strip_leading_zeros else digits
        lead = digits[:len(digits) - len(core)]
        head, tail = (core[:group_digits], core[group_digits:]) if group_digits and len(core) > group_digits else (core, "")
        keep, rest = head[:preserve_prefix], head[preserve_prefix:]
        if not rest:
            return v

        def make(attempt):
            d = self._keyed(domain, head, attempt)
            out = [d[i % len(d)] % 10 for i in range(len(rest))]
            if nonzero_lead and not keep:
                out[0] = 1 + d[0] % 9
            return keep + "".join(map(str, out))
        token = self._injective(domain, head, make)
        if tail and not keep_suffix:
            sub = f"{domain}\x1d{head}"
            tail = self._injective(sub, tail, lambda a: "".join(
                str(b % 10) for b in self._keyed(sub, tail, a)[:len(tail)]))
        new = iter(lead + token + tail)
        return "".join(next(new) if ch.isdigit() else ch for ch in v)

    def token_alnum(self, v, domain):
        # Keyed on letters and digits only, so "M1234" and "M1234-" (compound-meter low-flow
        # register) keep the same core token and the trailing dash survives.
        key = re.sub(r"[^A-Za-z0-9]", "", v).upper()
        if not key:
            return v

        def make(attempt):
            d = self._keyed(domain, key, attempt)
            return "".join(str(d[i % len(d)] % 10) if ch.isdigit() else string.ascii_uppercase[d[i % len(d)] % 26]
                           for i, ch in enumerate(key))
        new = iter(self._injective(domain, key, make))
        out = []
        for ch in v:
            if ch.isalnum():
                c = next(new)
                out.append(c.lower() if ch.islower() else c)
            else:
                out.append(ch)
        return "".join(out)

    def pseudo_word(self, word, domain):
        """Same-length pronounceable stand-in: vowels -> vowels, consonants -> consonants, case kept."""
        key = word.upper()

        def make(attempt):
            d = self._keyed(domain, key, attempt)
            return "".join(VOWELS[d[i % len(d)] % 5] if ch in VOWELS else
                           (CONSONANTS[d[i % len(d)] % 21] if ch.isalpha() else ch)
                           for i, ch in enumerate(key))
        token = self._injective(domain, key, make)
        return "".join(t.lower() if o.islower() else t for t, o in zip(token, word))

    def pseudo_text(self, v, domain, keep_words, number_domain=None):
        out = []
        for tok in re.findall(r"[A-Za-z]+|\d+|[^A-Za-z\d]+", v):
            if tok.isalpha():
                out.append(tok if len(tok) == 1 or tok.upper() in keep_words else self.pseudo_word(tok, domain))
            elif tok.isdigit():
                out.append(self.token_digits(tok, number_domain) if number_domain else tok)
            else:
                out.append(tok)
        return "".join(out)

    def fake_name(self, v, domain):
        f = self.faker_for(domain, v.strip().casefold())
        first, last = f.first_name(), f.last_name()
        out = f"{last}, {first}" if "," in v else f"{first} {last}"
        return self.match_case(v, out)

    def fake_company(self, v, domain):
        return self.match_case(v, self.faker_for(domain, v.strip().casefold()).company())

    def fake_street(self, v, domain):
        r, f = self.rng(domain, v), self.faker_for(domain, v.strip().casefold())
        return self.match_case(v, f"{r.randint(1, 9999)} {f.street_name()}")

    def fake_email(self, v, domain):
        return f"user{self.digest(domain, v.strip().casefold(), 5).hex()}@example.com"

    def hash_token(self, v, domain):
        return "TKN" + self.digest(domain, v, 6).hex().upper()

    def shift_date(self, v, domain, max_days=180, mode="global", formats=None):
        for fmt in formats or DATE_FORMATS:
            try:
                dt = datetime.strptime(v.strip(), fmt)
            except ValueError:
                continue
            key = "__global__" if mode == "global" else v
            days = self.rng(domain + ":shift", key).randint(-max_days, max_days) or 1
            return (dt + timedelta(days=days)).strftime(fmt), True
        return v, False

    def analyzer(self):
        if self._analyzer is None:
            from presidio_analyzer import AnalyzerEngine  # loads spaCy en_core_web_lg
            self._analyzer = AnalyzerEngine()
        return self._analyzer

    def redact_text(self, v, domain, entities=None, threshold=0.5):
        results = self.analyzer().analyze(text=v, language="en", entities=entities or PRESIDIO_ENTITIES,
                                          score_threshold=threshold)
        out = v
        for r in sorted(results, key=lambda x: x.start, reverse=True):
            frag = v[r.start:r.end]
            out = out[:r.start] + f"<{r.entity_type}_{self.digest(domain + ':' + r.entity_type, frag.casefold(), 2).hex()}>" + out[r.end:]
        return out


def apply_action(p: Pseudonymizer, spec: dict, value: str, column: str, stats: dict):
    action = spec.get("action", "keep")
    domain = spec.get("domain") or column.strip().lower()
    key = (action, domain, value, json.dumps(spec, sort_keys=True))
    if key in p.cache:
        return p.cache[key]
    if action == "keep":
        out = value
    elif action in ("blank", "drop"):
        out = ""
    elif action == "token_digits":
        out = p.token_digits(value, domain, int(spec.get("preserve_prefix", 0)), int(spec.get("group_digits", 0)),
                             bool(spec.get("keep_suffix", False)), bool(spec.get("strip_leading_zeros", False)),
                             bool(spec.get("nonzero_lead", False)))
    elif action == "fake_phone":
        out = p.token_digits(value, domain, int(spec.get("preserve_prefix", 3)))
    elif action == "token_alnum":
        out = p.token_alnum(value, domain)
    elif action == "fake_name":
        out = p.fake_name(value, domain)
    elif action == "fake_company":
        out = p.fake_company(value, domain)
    elif action == "fake_street":
        out = p.fake_street(value, domain)
    elif action == "fake_email":
        out = p.fake_email(value, domain)
    elif action == "hash":
        out = p.hash_token(value, domain)
    elif action == "mask_last":
        n = int(spec.get("keep", 4))
        out = re.sub(r"[A-Za-z0-9]", "X", value[:-n]) + value[-n:] if len(value) > n else value
    elif action == "shift_date":
        out, ok = p.shift_date(value, domain, int(spec.get("max_days", 180)), spec.get("mode", "global"), spec.get("formats"))
        if not ok:
            stats["unparseable"] = stats.get("unparseable", 0) + 1
            if spec.get("on_unparseable", "keep") == "blank":
                out = ""
    elif action == "pseudo_name":
        # digits inside name columns (phone numbers or IDs typed into a contact field) are tokenized too
        out = p.pseudo_text(value, domain, NAME_KEEP, number_domain=domain + ":number")
    elif action == "pseudo_address":
        out = p.pseudo_text(value, domain, ADDRESS_KEEP, number_domain=domain + ":number")
    elif action == "round":
        try:
            out = f"{round(float(value), int(spec.get('decimals', 2))):.{int(spec.get('decimals', 2))}f}"
        except ValueError:
            out = value
            stats["unparseable"] = stats.get("unparseable", 0) + 1
    elif action == "redact_text":
        out = p.redact_text(value, domain, spec.get("entities"), float(spec.get("threshold", 0.5)))
    elif action == "token_embedded_numbers":
        # a mostly-code column with an identifier embedded in some cells, e.g. a printed report that put the
        # meter number in the rate-table column ("WB2   12345678"): tokenize each word with min_digits+ digits
        # exactly as the identifier's own column is tokenized (same domain and options), keep everything else
        n = int(spec.get("min_digits", 5))
        out = re.sub(r"[A-Za-z0-9]*\d[A-Za-z0-9]*",
                     lambda m: p.token_digits(m.group(0), domain, int(spec.get("preserve_prefix", 0)), int(spec.get("group_digits", 0)),
                                              bool(spec.get("keep_suffix", False)), bool(spec.get("strip_leading_zeros", False)),
                                              bool(spec.get("nonzero_lead", False)))
                     if sum(ch.isdigit() for ch in m.group(0)) >= n else m.group(0), value)
    else:
        raise SystemExit(f"unknown action '{action}' for column '{column}'")
    p.cache[key] = out
    return out


# ---------------------------------------------------------------- io
def list_inputs(paths):
    for p in paths:
        p = Path(p)
        if p.is_dir():
            for f in sorted(p.iterdir()):
                if f.suffix.lower() in (".csv", ".xlsx", ".xlsm", ".txt") and f.is_file():
                    yield f
        elif p.is_file():
            yield p


def _is_date_format(fmt: str) -> bool:
    if not fmt or fmt in ("General", "@"):
        return False
    f = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', "", fmt).lower()   # drop literals, colours, escapes
    return any(t in f for t in ("y", "d")) or "mmm" in f


def _column_formats(ws, col, first_row, last_row, out):
    """Fill out[row] = NumberFormat for rows first_row..last_row, splitting only where formats differ."""
    rng = ws.Range(ws.Cells(first_row, col), ws.Cells(last_row, col))
    fmt = rng.NumberFormat
    if fmt is not None or first_row == last_row:
        for r in range(first_row, last_row + 1):
            out[r] = fmt or "General"
        return
    mid = (first_row + last_row) // 2
    _column_formats(ws, col, first_row, mid, out)
    _column_formats(ws, col, mid + 1, last_row, out)


def _display_text(v, fmt):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        if _is_date_format(fmt):
            seconds = round(float(v) * 86400)            # Excel stores date-times as fractional days
            dt = datetime(1899, 12, 30) + timedelta(seconds=seconds)
            has_time = "h" in fmt.lower() or seconds % 86400 != 0
            return dt.strftime("%m/%d/%Y %H:%M" if has_time else "%m/%d/%Y")
        if float(v).is_integer():
            return str(int(v))
        return repr(float(v))
    return str(v)


def read_protected_xlsx(path: Path, sheets="all"):
    """Yield (sheet_name, DataFrame of strings) from a rights-protected workbook via Excel COM.

    Values are read in memory under the user's own rights; no unprotected copy is written.
    Numbers and dates are turned back into the text Excel would display (dates as MM/DD/YYYY).
    Row and column positions are kept, so report layouts survive.
    """
    import gc
    import pythoncom
    import win32com.client
    import win32process
    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")   # private instance, never the user's Excel
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        _thread, excel_pid = win32process.GetWindowThreadProcessId(xl.Hwnd)
    except Exception:
        excel_pid = None
    wb = ws = used = None
    try:
        wb = xl.Workbooks.Open(str(path), 0, True)
        names = [ws.Name for ws in wb.Worksheets]
        if sheets != "all":
            wanted = sheets if isinstance(sheets, list) else [sheets]
            names = [names[s] if isinstance(s, int) else s for s in wanted]
        for name in names:
            ws = wb.Worksheets(name)
            used = ws.UsedRange
            values = used.Value2
            if values is None:
                yield name, pd.DataFrame()
                continue
            if not isinstance(values, tuple):
                values = ((values,),)
            r0, c0 = used.Row, used.Column
            n_rows, n_cols = len(values), len(values[0])
            grid = [[""] * (c0 - 1 + n_cols) for _ in range(r0 - 1 + n_rows)]
            for j in range(n_cols):
                if not any(isinstance(values[i][j], (int, float)) and not isinstance(values[i][j], bool)
                           for i in range(n_rows)):
                    for i in range(n_rows):
                        grid[r0 - 1 + i][c0 - 1 + j] = _display_text(values[i][j], "General")
                    continue
                fmts = {}
                _column_formats(ws, c0 + j, r0, r0 + n_rows - 1, fmts)
                for i in range(n_rows):
                    grid[r0 - 1 + i][c0 - 1 + j] = _display_text(values[i][j], fmts.get(r0 + i, "General"))
            yield name, pd.DataFrame(grid, dtype=str)
    finally:
        ws = used = None
        if wb is not None:
            wb.Close(False)
            wb = None
        xl.Quit()
        xl = None
        gc.collect()
        pythoncom.CoUninitialize()
        if excel_pid:                      # make sure the private Excel instance is gone
            import subprocess
            for _ in range(20):
                alive = subprocess.run(["tasklist", "/FI", f"PID eq {excel_pid}", "/NH"],
                                       capture_output=True, text=True).stdout
                if str(excel_pid) not in alive:
                    break
                time.sleep(0.25)
            else:
                subprocess.run(["taskkill", "/PID", str(excel_pid), "/F"], capture_output=True)


def read_grids(path: Path, sheets="all"):
    """Yield (sheet_name, DataFrame of strings, no header)."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import zipfile
        if not zipfile.is_zipfile(path):    # rights-protected (encrypted container): read through Excel
            yield from read_protected_xlsx(path, sheets)
            return
    if path.suffix.lower() in (".csv", ".txt"):
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                df = pd.read_csv(path, dtype=str, header=None, keep_default_na=False, encoding=enc,
                                 sep=None if path.suffix.lower() == ".txt" else ",", engine="python" if path.suffix.lower() == ".txt" else "c")
                yield None, df
                return
            except UnicodeDecodeError:
                continue
        raise SystemExit(f"can't decode {path.name}")
    xl = pd.ExcelFile(path)
    names = xl.sheet_names
    if sheets != "all":
        wanted = sheets if isinstance(sheets, list) else [sheets]
        names = [names[s] if isinstance(s, int) else s for s in wanted]
    for name in names:
        yield name, xl.parse(name, header=None, dtype=str, keep_default_na=False)


def guess_header_row(df: pd.DataFrame, max_rows=30):
    best, best_score = 0, -1
    for i in range(min(max_rows, len(df))):
        row = [str(x).strip() for x in df.iloc[i].tolist()]
        filled = [x for x in row if x]
        if not filled:
            continue
        labels = [x for x in filled if len(x) <= 40 and not re.fullmatch(r"[-+\d.,/: $%]+", x)]
        score = len(labels) if len(filled) >= 0.5 * len(row) else -1
        if score > best_score:
            best, best_score = i, score
        if score >= 0.8 * len(row) and score > 2:
            return i
    return best


# ---------------------------------------------------------------- commands
def cmd_suggest(a):
    from presidio_analyzer import AnalyzerEngine
    analyzer = AnalyzerEngine()
    policy = {"version": 1,
              "defaults": {"skip_values": SKIP_DEFAULT, "unknown_columns": "stop"},
              "files": []}
    print(f"{'file':40} {'sheet':10} {'column':30} suggestion   evidence")
    for f in list_inputs(a.inputs):
        for sheet, df in read_grids(f, "all"):
            h = guess_header_row(df)
            headers = [str(x).strip() for x in df.iloc[h].tolist()] if len(df) else []
            body = df.iloc[h + 1:]
            cols = {}
            for ci, name in enumerate(headers):
                if not name:
                    continue
                spec, why = {"action": "keep"}, "no rule matched"
                for rx, s, reason in NAME_RULES:
                    if re.search(rx, name, re.I):
                        spec, why = dict(s), f"name looks like {reason}"
                        break
                values = [v for v in body.iloc[:, ci].tolist() if str(v).strip() and not any(re.match(p, str(v)) for p in SKIP_DEFAULT)]
                sample = random.Random(0).sample(values, min(a.sample, len(values))) if values else []
                found = {}
                for v in sample:
                    for r in analyzer.analyze(text=str(v)[:500], language="en", entities=PRESIDIO_ENTITIES[:-1], score_threshold=0.6):
                        found[r.entity_type] = found.get(r.entity_type, 0) + 1
                if spec["action"] == "keep" and sample:
                    lens = sorted(len(str(v)) for v in sample)
                    looks_like_code = lens[len(lens) // 2] <= 6 or len(set(sample)) / len(sample) < 0.3
                    hits = {k: n for k, n in found.items()
                            if n >= max(2, 0.2 * len(sample)) and (k in STRONG_ENTITIES or not looks_like_code)}
                    if hits:
                        spec, why = {"action": "redact_text"}, "presidio found " + ", ".join(f"{k} {n}/{len(sample)}" for k, n in hits.items())
                    elif looks_like_code:
                        why = "looks like a code column (short or repetitive values)"
                evidence = why + (f"; presidio: {found}" if found and "presidio" not in why else "")
                cols[name] = spec
                print(f"{f.name[:40]:40} {str(sheet)[:10]:10} {name[:30]:30} {spec['action']:12} {evidence}")
            policy["files"].append({"match": f.name, "sheets": "all" if sheet is not None else None,
                                    "header_row": int(h), "columns": cols})
    for entry in policy["files"]:
        if entry["sheets"] is None:
            del entry["sheets"]
    Path(a.out).write_text("# DRAFT policy generated by deidentify.py suggest. REVIEW every column before use.\n"
                           + yaml.safe_dump(policy, sort_keys=False, allow_unicode=True), encoding="utf-8")
    print(f"\nDraft written to {a.out}. Review it: header rows, join-key domains, and every 'keep'.")


def cmd_apply(a):
    policy = yaml.safe_load(Path(a.policy).read_text(encoding="utf-8"))
    defaults = policy.get("defaults", {})
    skip_global = [re.compile(x) for x in defaults.get("skip_values", SKIP_DEFAULT)]
    unknown = defaults.get("unknown_columns", "stop")
    out_dir = Path(a.out)
    in_dirs = [Path(x).resolve() for x in a.inputs if Path(x).is_dir()]
    if any(out_dir.resolve() == d or d in out_dir.resolve().parents for d in in_dirs):
        sys.exit("Refusing to write inside the input folder. Choose a separate output folder.")
    if "onedrive" in str(out_dir.resolve()).lower():
        print("WARNING: the output folder is inside OneDrive and will sync. Make sure that's approved for this data.")
    out_dir.mkdir(parents=True, exist_ok=True)
    p = Pseudonymizer(load_salt())
    report = {"policy": str(a.policy), "started": time.strftime("%Y-%m-%d %H:%M:%S"), "files": []}
    warnings = []

    for f in list_inputs(a.inputs):
        entry = next((e for e in policy.get("files", []) if fnmatch.fnmatch(f.name, e["match"])), None)
        if entry is None:
            print(f"SKIP {f.name}: no policy entry (nothing written for this file)")
            report["files"].append({"file": f.name, "status": "skipped: no policy"})
            continue
        sheets_out, file_rep = {}, {"file": f.name, "sheets": []}
        for sheet, df in read_grids(f, entry.get("sheets", "all")):
            h = int(entry.get("header_row", 0))
            if df.empty:
                sheets_out[sheet] = df
                file_rep["sheets"].append({"sheet": sheet, "rows": 0})
                continue
            headers = [str(x).strip() for x in df.iloc[h].tolist()]
            # Data under a blank header would otherwise pass through untouched: name it so the policy must decide.
            for ci, name in enumerate(headers):
                if not name and any(str(x).strip() for x in df.iloc[h + 1:, ci].tolist()):
                    headers[ci] = f"(column {ci + 1})"
            header_set = {x for x in headers if x}
            cols = entry.get("columns", {})
            missing = [c for c in headers if c and c not in cols]
            if missing and unknown == "stop":
                sys.exit(f"{f.name} [{sheet}]: columns not in the policy: {missing}. Add them, or set "
                         f"defaults.unknown_columns to keep/blank/redact_text explicitly.")
            sheet_rep = {"sheet": sheet, "rows": int(len(df)), "header_row": h, "columns": {}}
            for ci, name in enumerate(headers):
                if not name:
                    continue
                spec = cols.get(name) or {"action": unknown if unknown != "stop" else "keep"}
                if spec.get("action", "keep") == "keep":
                    # safeguard: a kept column holding long numbers may be carrying identifiers (account, meter or
                    # receipt numbers) that a report layout pushed into it; review it unless the policy says it's fine
                    long_cells = sum(1 for v in df.iloc[h + 1:, ci].tolist() if LONG_NUMBER_RE.search(str(v)))
                    sheet_rep["columns"][name] = {"action": "keep"}
                    if long_cells and not spec.get("long_numbers_ok"):
                        sheet_rep["columns"][name]["cells_with_long_numbers"] = long_cells
                        warnings.append(f"{f.name} [{sheet}] keep column '{name}': {long_cells} cell(s) contain 7+ digit numbers; "
                                        f"check they aren't identifiers, then set long_numbers_ok: true or change the action")
                    continue
                skips = skip_global + [re.compile(x) for x in spec.get("skip_values", [])]
                stats = {"action": spec.get("action"), "changed": 0, "skipped": 0}
                col = df.iloc[:, ci].tolist()
                for ri in range(h + 1, len(col)):
                    v = str(col[ri])
                    if v.strip() in header_set or any(rx.match(v) for rx in skips):   # repeated report headers
                        stats["skipped"] += 1
                        continue
                    nv = apply_action(p, spec, v, name, stats)
                    if nv != v:
                        stats["changed"] += 1
                    col[ri] = nv
                df.iloc[:, ci] = col
                sheet_rep["columns"][name] = stats
            sheets_out[sheet] = df
            file_rep["sheets"].append(sheet_rep)
        target = out_dir / f.name
        if f.suffix.lower() in (".csv", ".txt"):
            next(iter(sheets_out.values())).to_csv(target, header=False, index=False, encoding="utf-8")
        else:
            # xlsx can't store control characters (report exports carry form feeds at page breaks).
            # Map each one to its Unicode "control picture" (e.g. \x0c -> U+240C), which is reversible
            # and keeps page breaks visible to the reshape step.
            from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
            replaced = 0
            for sheet, df in sheets_out.items():
                def fix(s):
                    nonlocal replaced
                    if isinstance(s, str) and ILLEGAL_CHARACTERS_RE.search(s):
                        replaced += 1
                        return ILLEGAL_CHARACTERS_RE.sub(lambda mm: chr(0x2400 + ord(mm.group(0))), s)
                    return s
                sheets_out[sheet] = df.apply(lambda c: c.map(fix))
            file_rep["control_characters_mapped_cells"] = replaced
            # Write every cell as text: openpyxl would otherwise turn report lines such as "=========" into formulas.
            from openpyxl import Workbook
            wb_out = Workbook()
            wb_out.remove(wb_out.active)
            for sheet, df in sheets_out.items():
                ws_out = wb_out.create_sheet(title=str(sheet)[:31])
                for r_i, row in enumerate(df.itertuples(index=False, name=None), start=1):
                    for c_i, val in enumerate(row, start=1):
                        if val is None or (isinstance(val, float) and pd.isna(val)) or val == "":
                            continue
                        cell = ws_out.cell(row=r_i, column=c_i, value=str(val))
                        cell.data_type = "s"
            wb_out.save(target.with_suffix(".xlsx"))
        print(f"OK   {f.name}")
        report["files"].append(file_rep)

    report["finished"] = time.strftime("%Y-%m-%d %H:%M:%S")
    report["distinct_tokens_by_domain"] = {d: len(m) for d, m in sorted(p.fwd.items())}
    report["collision_fallbacks_by_domain"] = dict(p.collision_fallbacks)   # should be empty
    report["warnings"] = warnings
    rp = out_dir / "deidentify-report.json"
    rp.write_text(json.dumps(report, indent=1, default=str), encoding="utf-8")
    print(f"Report (counts only): {rp}")
    for w in warnings:
        print("WARNING:", w)
    if p.collision_fallbacks:
        print(f"WARNING: token space exhausted in {p.collision_fallbacks}; joins in those domains may merge values.")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("suggest", help="draft a policy (prints column names and suggestions only)")
    s.add_argument("--in", dest="inputs", nargs="+", required=True)
    s.add_argument("--out", required=True)
    s.add_argument("--sample", type=int, default=50, help="values per column scanned with Presidio")
    s.set_defaults(func=cmd_suggest)
    p = sub.add_parser("apply", help="apply a reviewed policy")
    p.add_argument("--policy", required=True)
    p.add_argument("--in", dest="inputs", nargs="+", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(func=cmd_apply)
    a = ap.parse_args(argv)
    a.func(a)
    return 0


if __name__ == "__main__":
    sys.exit(main())
