"""Read client extracts as text, exactly as received.

Every cell becomes a string: blanks are '' (never None), whole numbers have no '.0', dates are
MM/DD/YYYY. Nothing is trimmed or "fixed" here; that's the reshape step's job, and it's logged.
"""
from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, dt.datetime):
        return v.strftime("%m/%d/%Y") if v.time() == dt.time(0) else v.strftime("%m/%d/%Y %H:%M:%S")
    if isinstance(v, dt.date):
        return v.strftime("%m/%d/%Y")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M:%S")
    return str(v)


def list_sheets(path: Path) -> list[tuple[str, int, int]]:
    """(sheet name, rows, columns) for every sheet, so a run report can show what was NOT used."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            n = sum(1 for _ in f)
        return [("", n, 0)]
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return [(ws.title, ws.max_row or 0, ws.max_column or 0) for ws in wb.worksheets]
    finally:
        wb.close()


def read_grid(path: Path, sheet: str | int | None = None, delimiter: str = ",", encoding: str = "utf-8-sig") -> list[list[str]]:
    """All rows of one sheet (or the CSV) as lists of strings, header row included."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding=encoding) as f:
            return [list(r) for r in csv.reader(f, delimiter=delimiter)]
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None:
            ws = wb.worksheets[0]
        elif isinstance(sheet, int):
            ws = wb.worksheets[sheet]
        else:
            ws = wb[sheet]
        rows = [[cell_text(v) for v in r] for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows]
