
"""Generate the familiar Field Info workbook from a metadata-only mapping draft."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    "Required", "Friendly name/field", "Description", "GP field", "GP type", "Length", "Source system",
    "Source table/file", "Source field", "Cleansing rules", "Transformation rules", "Remarks",
    "Agent confidence", "Open question",
]


def _safe_sheet(name: str) -> str:
    return name[:31].replace("/", "-").replace("\\", "-").replace("?", "")


def write_workbook(draft: dict | str | Path, out_path: str | Path) -> Path:
    if not isinstance(draft, dict):
        draft = json.loads(Path(draft).read_text(encoding="utf-8"))
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    questions = []
    for module in draft["modules"]:
        ws = wb.create_sheet(_safe_sheet(module))
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for r in [x for x in draft["mappings"] if x["module"] == module]:
            t = r["target"]
            open_q = r.get("open_question") or ""
            if open_q:
                questions.append([module, t.get("field"), open_q, r.get("confidence")])
            ws.append([
                "Y" if t.get("required") else "N",
                t.get("field") or "",
                t.get("description") or "",
                t.get("gp_field") or "",
                t.get("gp_type") or "",
                t.get("length") or "",
                "SAMPLE extract profile" if r.get("source_column") else "",
                r.get("source_file") or "",
                r.get("source_column") or "",
                "",
                r.get("transformation") or "",
                f"Agent derivation: {r.get('derivation')}",
                r.get("confidence"),
                open_q,
            ])
        _tidy(ws, border)
    oq = wb.create_sheet("Open questions")
    oq.append(["Module", "Field", "Open question", "Agent confidence"])
    for cell in oq[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = border
    for q in questions:
        oq.append(q)
    _tidy(oq, border)
    wb.save(out_path)
    return out_path


def _tidy(ws, border):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        1: 10, 2: 24, 3: 44, 4: 18, 5: 14, 6: 10, 7: 22, 8: 32, 9: 24, 10: 28, 11: 34, 12: 34, 13: 16, 14: 44,
    }
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    ws.row_dimensions[1].height = 32
