"""Generate the per-take data integrity report for process P14.

The report is deterministic: all counts come from reconciliation views, run logs, answer-key rows,
runbook execution rows, specs and the approved/fallback triage proposal file. It emits no record-level
values: only module names, rule codes, counts, control totals, Jira references and recommendations.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import html
import json
import os
import subprocess
from collections import defaultdict
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.shared import Inches

from . import sqlio
from .triage import MODULE_NAMES

MODULE_ORDER = ["UMRM0101", "UMRM0111", "UM00600", "UM00600H", "UM00400", "UM00300", "UM20400", "UM01900", "UM30300"]
GUIDELINES = [
    "Correct data in the legacy source whenever possible so every later take receives the same fix.",
    "Keep account, customer, location, equipment and meter identifiers consistent across every extract.",
    "Provide dates in an unambiguous format and avoid free-text dates.",
    "Use configured CSM codes for service types, rates, tariffs, classes and deposit types.",
    "Confirm that excluded service codes are intentional and documented.",
    "Resolve duplicate keys before load: customer, location, equipment, connection and read keys must be unique.",
    "Ensure each connection that should load has matching equipment when the service is metered.",
    "Keep address lines within target field lengths and review postal-standardization misses.",
    "Use two-character state/province abbreviations for licence and address state fields.",
    "Do not place meter identifiers or licence numbers in unrelated free-text fields.",
    "Check beginning-balance and deposit amounts against finance control totals.",
    "Review inactive/finaled accounts that still carry balances, deposits or meter history.",
    "Preserve source extracts exactly as submitted; changes after submission should be made in the next take.",
    "Treat warnings as loadable records that still require a fix-in-source decision.",
    "Record every business decision in the spec or gate history so the next take repeats it automatically.",
]


def q(cn, sql: str, *params) -> list[dict[str, Any]]:
    return sqlio.query(cn, sql, *params)


def money(v: Any) -> str:
    if v is None:
        return ""
    return f"{float(v):,.2f}"


def num(v: Any) -> str:
    if v is None:
        return ""
    return f"{int(v):,}"


def display_take_date(value: Any) -> str:
    text = str(value or "")
    if len(text) == 8 and text.isdigit():
        return f"{text[4:8]}-{text[0:2]}-{text[2:4]}"
    return text


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:12]


def read_yaml(path: Path) -> dict[str, Any]:
    return yaml.safe_load(path.read_text(encoding="utf-8")) or {}


def spec_hashes(client_dir: Path) -> list[tuple[str, str]]:
    files = [client_dir / "take.yaml", client_dir / "runbook.yaml", client_dir / "jira-refs.yaml"]
    files += sorted((client_dir / "reshape").glob("*.yaml"))
    files += sorted((client_dir / "mapping").glob("*.yaml"))
    return [(str(p.relative_to(client_dir)), sha(p)) for p in files if p.exists()]


def conversion_plan(client_dir: Path) -> list[dict[str, str]]:
    reshapes = []
    for p in sorted((client_dir / "reshape").glob("*.yaml")):
        s = read_yaml(p)
        if s.get("kind") == "reshape":
            source = (s.get("source") or {}).get("file") or s.get("source") or p.name
            reshapes.append({"spec": p.name, "file": str(source), "node": str(s.get("target", "")), "map": "reshape"})
    mappings = []
    for p in sorted((client_dir / "mapping").glob("*.yaml")):
        s = read_yaml(p)
        if s.get("kind") == "mapping":
            mappings.append({"spec": p.name, "file": str(s.get("source", "")), "node": str(s.get("target", s.get("module", ""))), "map": str(s.get("module", p.stem))})
    return reshapes + mappings


def load_jira_refs(client_dir: Path) -> dict[str, str]:
    p = client_dir / "jira-refs.yaml"
    data = read_yaml(p)
    out = {}
    for rule, refs in (data.get("refs") or {}).items():
        out[rule] = ", ".join(refs or [])
    return out


def load_triage(work: Path, take: str) -> dict[str, dict[str, str]]:
    candidates = [work / "triage" / "triage-proposals.json", Path.home() / "CogsdaleWork" / "sample-run" / take / "triage" / "triage-proposals.json"]
    path = next((p for p in candidates if p.exists()), None)
    if path is None:
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    props = data.get("proposals", data if isinstance(data, list) else [])
    by_rule: dict[str, dict[str, Any]] = {}
    for p in props:
        rules = str(p.get("RuleSet", "")).split("+")
        for rule in rules:
            if not rule:
                continue
            row = by_rule.setdefault(rule, {"records": 0, "owners": defaultdict(int), "actions": defaultdict(int), "titles": defaultdict(int)})
            row["records"] += int(p.get("Records") or 0)
            row["owners"][p.get("Owner") or p.get("FinalOwner") or ""] += int(p.get("Records") or 0)
            row["actions"][p.get("Action") or ""] += int(p.get("Records") or 0)
            row["titles"][p.get("Title") or ""] += int(p.get("Records") or 0)
    result = {}
    for rule, row in by_rule.items():
        owner = max(row["owners"], key=row["owners"].get) if row["owners"] else ""
        action = max(row["actions"], key=row["actions"].get) if row["actions"] else ""
        title = max(row["titles"], key=row["titles"].get) if row["titles"] else ""
        result[rule] = {"owner": owner, "recommendation": action, "title": title}
    return result


def collect(cn, client: str, take: str, client_dir: Path, work: Path) -> dict[str, Any]:
    data: dict[str, Any] = {"client": client, "take": take, "generated_at": dt.datetime.now(dt.timezone.utc)}
    from .specs import ClientSpecs
    data["take_yaml"] = ClientSpecs(client_dir, take).take        # the selected take's own variables (date, files, notes)
    data["runbook_yaml"] = read_yaml(client_dir / "runbook.yaml")
    data["spec_hashes"] = spec_hashes(client_dir)
    data["plan"] = conversion_plan(client_dir)
    data["jira"] = load_jira_refs(client_dir)
    data["triage"] = load_triage(work, take)
    data["modules"] = q(cn, "SELECT * FROM recon.v_ModuleSummary WHERE ClientId = ? AND TakeId = ?", client, take)
    data["issues"] = q(cn, "SELECT * FROM recon.v_IssueSummary WHERE ClientId = ? AND TakeId = ?", client, take)
    data["totals"] = q(cn, "SELECT * FROM recon.v_ControlTotal WHERE ClientId = ? AND TakeId = ?", client, take)
    data["answer"] = q(cn, "SELECT Module, Metric, Expected, Actual, Status, Reason, ApprovedBy FROM run.AnswerKeyCheck WHERE ClientId = ? AND TakeId = ?", client, take)
    data["answer_summary"] = q(cn, "SELECT Status, COUNT(*) AS Cnt FROM run.AnswerKeyCheck WHERE ClientId = ? AND TakeId = ? GROUP BY Status", client, take)
    data["runbook"] = q(cn, "SELECT s.StepNo, s.Block, s.Title, s.Kind, s.Executor, s.PrototypeMode, r.Status, r.Evidence, r.Message FROM meta.RunbookStep s LEFT JOIN run.RunbookStepRun r ON r.ClientId = s.ClientId AND r.TakeId = ? AND r.StepNo = s.StepNo WHERE s.ClientId = ? ORDER BY s.StepNo", take, client)
    data["runlog"] = q(cn, "SELECT StepName, StartedAt, EndedAt, RowsAffected, Status FROM run.RunLog WHERE RunLogId IN (SELECT MAX(RunLogId) FROM run.RunLog WHERE ClientId = ? AND TakeId = ? GROUP BY StepName) ORDER BY StepName", client, take)
    results_path = work / "run-results.json"
    answer_path = work / "answer-key-check.json"
    data["run_results"] = json.loads(results_path.read_text(encoding="utf-8")) if results_path.exists() else None
    data["answer_file"] = json.loads(answer_path.read_text(encoding="utf-8")) if answer_path.exists() else None
    return data


def verify_numbers(data: dict[str, Any]) -> list[str]:
    checks = []
    rr = data.get("run_results")
    if rr:
        by_mod = {r["Module"]: r for r in data["modules"]}
        for mod, vals in rr.get("modules", {}).items():
            if mod in by_mod:
                for metric, col in [("received", "Received"), ("loaded", "Loaded"), ("failed", "Failed")]:
                    if int(vals.get(metric, 0)) != int(by_mod[mod][col]):
                        raise ValueError(f"number mismatch {mod}.{metric}: JSON {vals.get(metric)} SQL {by_mod[mod][col]}")
        checks.append(f"Module counts match {len(by_mod)} SQL reconciliation rows and run-results.json.")
    af = data.get("answer_file")
    if af:
        actual = {r["Status"].lower(): int(r["Cnt"]) for r in data["answer_summary"]}
        for key, value in af.get("summary", {}).items():
            if int(value) != int(actual.get(key.lower(), 0)):
                raise ValueError(f"answer-key mismatch {key}: JSON {value} SQL {actual.get(key.lower(), 0)}")
        checks.append("Answer-key summary matches answer-key-check.json.")
    return checks


def h(text: Any) -> str:
    return html.escape("" if text is None else str(text))


def answer_summary_text(data: dict[str, Any]) -> str:
    if not data["take_yaml"].get("answer_key"):
        return "No answer key: this take has no known outcome to reproduce, so only the reconciliation checks apply"
    counts = {str(r["Status"]).upper(): int(r["Cnt"]) for r in data["answer_summary"]}
    order = ["PASS", "ALLOWED", "REVIEWED", "FAIL", "NEW"]
    return ", ".join(f"{counts.get(status, 0)} {status}" for status in order)


def report_title(data: dict[str, Any]) -> str:
    return f"Data Integrity Report - {data['client'].upper()} {data['take']}" + (" (SIMULATED)" if data["take_yaml"].get("simulated") else "")


def simulated_banner(data: dict[str, Any]) -> str | None:
    """A synthetic take's report says so up front: it is never a client delivery."""
    if not data["take_yaml"].get("simulated"):
        return None
    return "SIMULATED take, not a client delivery. " + " ".join(str(data["take_yaml"].get("notes") or "").split())


def schedule_text(take_yaml: dict[str, Any]) -> str:
    return (f"Take {take_yaml.get('take_id')} uses source-file date {display_take_date(take_yaml.get('take_date'))} and conversion date "
            f"{take_yaml.get('conversion_date')}. Every take follows the same schedule pattern: mapping approval, extract submission, "
            "reshape/build/validate, triage, load decision, publish decision and gate 6 acceptance.")


def html_report(data: dict[str, Any], checks: list[str]) -> str:
    take_yaml = data["take_yaml"]
    module_rows = {r["Module"]: r for r in data["modules"]}
    issues_by_mod: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in data["issues"]:
        issues_by_mod[r["Module"]].append(r)
    totals_by_mod: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in data["totals"]:
        totals_by_mod[r["Module"]].append(r)
    rb_status = defaultdict(int)
    for r in data["runbook"]:
        rb_status[r.get("Status") or "not_run"] += 1

    parts = ["""<!doctype html><html><head><meta charset='utf-8'><title>Data Integrity Report</title>
<style>
@page { size: Letter; margin: 0.6in; } body { font-family: 'Segoe UI', Arial, sans-serif; color: #1f2937; line-height: 1.35; }
h1 { color: #0f4c81; font-size: 24px; } h2 { color: #0f4c81; border-bottom: 1px solid #cbd5e1; padding-bottom: 3px; page-break-after: avoid; }
h3 { color: #334155; page-break-after: avoid; } table { border-collapse: collapse; width: 100%; margin: 8px 0 16px; font-size: 9.5pt; }
th, td { border: 1px solid #cbd5e1; padding: 4px 6px; vertical-align: top; } th { background: #e2e8f0; text-align: left; }
.badge { display: inline-block; border-radius: 10px; padding: 2px 8px; background: #e0f2fe; margin-right: 5px; }
.small { font-size: 9pt; color: #475569; } .pagebreak { page-break-before: always; } .right { text-align:right; }
</style></head><body>"""]
    parts.append(f"<h1>{h(report_title(data))}</h1>")
    if simulated_banner(data):
        parts.append(f"<p style='border:2px solid #b45309;background:#fef3c7;padding:8px 10px;font-weight:600'>{h(simulated_banner(data))}</p>")
    parts.append("<p class='small'>Generated from the conversion platform run. No record-level customer, account, address or meter identifiers are included.</p>")

    parts.append("<h2>1. Version control</h2><table><tr><th>Item</th><th>Value</th></tr>")
    parts.append(f"<tr><td>Take</td><td>{h(data['take'])}</td></tr><tr><td>Conversion date</td><td>{h(take_yaml.get('conversion_date'))}</td></tr><tr><td>Generated UTC</td><td>{h(data['generated_at'].isoformat(timespec='seconds'))}</td></tr>")
    parts.append(f"<tr><td>Answer-key reproduction</td><td>{h(answer_summary_text(data))}</td></tr></table>")
    parts.append("<table><tr><th>Spec file</th><th>SHA-256 prefix</th></tr>" + "".join(f"<tr><td>{h(p)}</td><td>{h(s)}</td></tr>" for p, s in data["spec_hashes"]) + "</table>")

    parts.append("<h2>2. General cleansing guidelines</h2><ol>" + "".join(f"<li>{h(x)}</li>" for x in GUIDELINES) + "</ol>")

    parts.append("<h2>3. Conversion planning</h2><table><tr><th>Source / spec</th><th>File or silver source</th><th>Node / target</th><th>Map</th></tr>")
    for row in data["plan"]:
        parts.append(f"<tr><td>{h(row['spec'])}</td><td>{h(row['file'])}</td><td>{h(row['node'])}</td><td>{h(row['map'])}</td></tr>")
    parts.append("</table>")

    parts.append("<h2>4. Schedule</h2>")
    parts.append(f"<p>{h(schedule_text(take_yaml))}</p>")

    parts.append("<h2>5. Conversion execution sequence</h2>")
    parts.append("<p>Prototype status summary: " + ", ".join(f"<span class='badge'>{h(k)}: {v}</span>" for k, v in sorted(rb_status.items())) + "</p>")
    parts.append("<table><tr><th>#</th><th>Block</th><th>Step</th><th>Kind</th><th>Executor</th><th>Status</th><th>Evidence / message</th></tr>")
    for r in data["runbook"]:
        msg = (r.get("Evidence") or "") + (" - " if r.get("Evidence") and r.get("Message") else "") + (r.get("Message") or "")
        parts.append(f"<tr><td class='right'>{r['StepNo']}</td><td>{h(r['Block'])}</td><td>{h(r['Title'])}</td><td>{h(r['Kind'])}</td><td>{h(r['Executor'])}</td><td>{h(r.get('Status') or 'not run')}</td><td>{h(msg)}</td></tr>")
    parts.append("</table>")

    parts.append("<h2>6. Processed items</h2>")
    for mod in MODULE_ORDER:
        if mod not in module_rows:
            continue
        m = module_rows[mod]
        parts.append(f"<h3>{h(mod)} - {h(MODULE_NAMES.get(mod, mod))}</h3>")
        parts.append(f"<p>Received <b>{num(m['Received'])}</b>; loaded <b>{num(m['Loaded'])}</b>; failed <b>{num(m['Failed'])}</b>.</p>")
        if issues_by_mod.get(mod):
            parts.append("<table><tr><th>Rule</th><th>Severity</th><th>Records</th><th>Expected</th><th>Actionable</th><th>Jira</th><th>Owner</th><th>Recommendation</th></tr>")
            for iss in sorted(issues_by_mod[mod], key=lambda x: x["RuleId"]):
                tri = data["triage"].get(iss["RuleId"], {})
                owner = tri.get("owner") or ("none" if int(iss["Actionable"]) == 0 else "client/cogsdale")
                rec = tri.get("recommendation") or ("No action; expected by approved exclusion." if int(iss["Actionable"]) == 0 else "Review and correct the source, mapping or setup before the next take.")
                parts.append(f"<tr><td>{h(iss['RuleId'])}</td><td>{h(iss['Severity'])}</td><td class='right'>{num(iss['Records'])}</td><td class='right'>{num(iss['Expected'])}</td><td class='right'>{num(iss['Actionable'])}</td><td>{h(data['jira'].get(iss['RuleId'], ''))}</td><td>{h(owner)}</td><td>{h(rec)}</td></tr>")
            parts.append("</table>")
        else:
            parts.append("<p>No validation issues were recorded for this module.</p>")
        if totals_by_mod.get(mod):
            parts.append("<table><tr><th>Control total</th><th>Records</th><th>Amount</th></tr>")
            for t in totals_by_mod[mod]:
                parts.append(f"<tr><td>{h(t['Name'])}</td><td class='right'>{num(t['Records'])}</td><td class='right'>{money(t['Amount'])}</td></tr>")
            parts.append("</table>")

    parts.append("<h2>7. Next steps / control tracking</h2>")
    actionable = [i for i in data["issues"] if int(i["Actionable"] or 0) > 0]
    if actionable:
        parts.append("<table><tr><th>Owner</th><th>Rule</th><th>Actionable records</th><th>Tracking</th></tr>")
        for iss in sorted(actionable, key=lambda x: (-int(x["Actionable"]), x["RuleId"])):
            tri = data["triage"].get(iss["RuleId"], {})
            parts.append(f"<tr><td>{h(tri.get('owner') or 'client/cogsdale')}</td><td>{h(iss['RuleId'])}</td><td class='right'>{num(iss['Actionable'])}</td><td>{h(data['jira'].get(iss['RuleId'], ''))}</td></tr>")
        parts.append("</table>")
    parts.append("<p>Pending allowances and reviewed differences remain visible in the answer-key appendix. Gate 6 acceptance is unsigned in this generated prototype report.</p>")

    parts.append("<h2>8. Data integrity acceptance</h2>")
    parts.append("<table><tr><th>Role</th><th>Name / signature</th><th>Date</th></tr><tr><td>Client project manager</td><td>Unsigned</td><td></td></tr><tr><td>Cogsdale project manager</td><td>Unsigned</td><td></td></tr><tr><td>Conversion lead</td><td>Unsigned</td><td></td></tr></table>")
    parts.append("<p>Acceptance is represented by gate 6. Until gate 6 is approved, this report is a draft evidence package.</p>")

    parts.append("<h2>Appendix A. Reproduction check</h2>")
    parts.append("<p>" + h(answer_summary_text(data)) + "</p>")
    parts.append("<table><tr><th>Module</th><th>Metric</th><th>Expected</th><th>Actual</th><th>Status</th></tr>")
    for r in data["answer"]:
        parts.append(f"<tr><td>{h(r['Module'])}</td><td>{h(r['Metric'])}</td><td>{h(r['Expected'])}</td><td>{h(r['Actual'])}</td><td>{h(r['Status'])}</td></tr>")
    parts.append("</table><h3>Verification checks</h3><ul>" + "".join(f"<li>{h(c)}</li>" for c in checks) + "</ul>")
    parts.append("</body></html>")
    return "\n".join(parts)


def write_docx(data: dict[str, Any], checks: list[str], path: Path) -> None:
    doc = Document()
    doc.add_heading(report_title(data), 0)
    if simulated_banner(data):
        doc.add_paragraph().add_run(simulated_banner(data)).bold = True
    doc.add_paragraph("Generated from the conversion platform run. No record-level identifiers are included.")
    doc.add_heading("1. Version control", 1)
    t = doc.add_table(rows=0, cols=2)
    for k, v in [("Take", data["take"]), ("Conversion date", data["take_yaml"].get("conversion_date")), ("Generated UTC", data["generated_at"].isoformat(timespec="seconds")), ("Answer-key reproduction", answer_summary_text(data))]:
        row = t.add_row().cells; row[0].text = str(k); row[1].text = str(v)
    doc.add_heading("2. General cleansing guidelines", 1)
    for g in GUIDELINES:
        doc.add_paragraph(g, style="List Number")
    doc.add_heading("3. Conversion planning", 1)
    table = doc.add_table(rows=1, cols=4); table.rows[0].cells[0].text="Source / spec"; table.rows[0].cells[1].text="File or silver source"; table.rows[0].cells[2].text="Node / target"; table.rows[0].cells[3].text="Map"
    for p in data["plan"]:
        row=table.add_row().cells; row[0].text=p["spec"]; row[1].text=p["file"]; row[2].text=p["node"]; row[3].text=p["map"]
    doc.add_heading("4. Schedule", 1)
    doc.add_paragraph(schedule_text(data["take_yaml"]))
    doc.add_heading("5. Conversion execution sequence", 1)
    table = doc.add_table(rows=1, cols=5); hdr=table.rows[0].cells
    for i, h0 in enumerate(["#", "Block", "Step", "Status", "Evidence"]): hdr[i].text=h0
    for r in data["runbook"]:
        row=table.add_row().cells; row[0].text=str(r["StepNo"]); row[1].text=str(r["Block"]); row[2].text=str(r["Title"]); row[3].text=str(r.get("Status") or "not run"); row[4].text=str(r.get("Evidence") or r.get("Message") or "")
    doc.add_heading("6. Processed items", 1)
    module_rows = {r["Module"]: r for r in data["modules"]}
    issues_by_mod=defaultdict(list)
    for r in data["issues"]: issues_by_mod[r["Module"]].append(r)
    for mod in MODULE_ORDER:
        if mod not in module_rows: continue
        m=module_rows[mod]; doc.add_heading(f"{mod} - {MODULE_NAMES.get(mod, mod)}", 2)
        doc.add_paragraph(f"Received {num(m['Received'])}; loaded {num(m['Loaded'])}; failed {num(m['Failed'])}.")
        if issues_by_mod.get(mod):
            tbl=doc.add_table(rows=1, cols=6); hdr=tbl.rows[0].cells
            for i,h0 in enumerate(["Rule", "Severity", "Records", "Expected", "Actionable", "Jira"]): hdr[i].text=h0
            for iss in sorted(issues_by_mod[mod], key=lambda x:x["RuleId"]):
                row=tbl.add_row().cells; row[0].text=iss["RuleId"]; row[1].text=iss["Severity"]; row[2].text=num(iss["Records"]); row[3].text=num(iss["Expected"]); row[4].text=num(iss["Actionable"]); row[5].text=data["jira"].get(iss["RuleId"], "")
    doc.add_heading("7. Next steps / control tracking", 1)
    doc.add_paragraph("Open fix-in-source items are listed by rule in the HTML version with owner and recommendation. Pending allowances remain in the reproduction appendix.")
    doc.add_heading("8. Data integrity acceptance", 1)
    acc=doc.add_table(rows=1, cols=3); acc.rows[0].cells[0].text="Role"; acc.rows[0].cells[1].text="Name / signature"; acc.rows[0].cells[2].text="Date"
    for role in ["Client project manager", "Cogsdale project manager", "Conversion lead"]:
        row=acc.add_row().cells; row[0].text=role; row[1].text="Unsigned"; row[2].text=""
    doc.add_heading("Appendix A. Reproduction check", 1)
    doc.add_paragraph(answer_summary_text(data))
    for c in checks: doc.add_paragraph(c, style="List Bullet")
    doc.sections[0].left_margin = Inches(0.6); doc.sections[0].right_margin = Inches(0.6)
    doc.save(path)


def render_pdf(html_path: Path, pdf_path: Path) -> None:
    script = pdf_path.with_suffix(".render.js")
    html_url = "file:///" + str(html_path.resolve()).replace("\\", "/")    # outside the f-string: Python 3.11 (Fabric) forbids backslashes in f-string expressions
    js = f"""
const path = require('path');
const playwright = require(path.join(process.env.APPDATA, 'npm/node_modules/@microsoft/powerbi-report-authoring-cli/node_modules/playwright-core'));
(async () => {{
  const browser = await playwright.chromium.launch({{executablePath: 'C:\\\\Program Files (x86)\\\\Microsoft\\\\Edge\\\\Application\\\\msedge.exe', headless: true}});
  const page = await browser.newPage();
  await page.goto({json.dumps(html_url)}, {{waitUntil: 'networkidle'}});
  await page.pdf({{path: {json.dumps(str(pdf_path.resolve()))}, format: 'Letter', printBackground: true, margin: {{top: '0.6in', bottom: '0.6in', left: '0.6in', right: '0.6in'}}}});
  await browser.close();
}})().catch(e => {{ console.error(e); process.exit(1); }});
"""
    script.write_text(js, encoding="utf-8")
    subprocess.run(["node", str(script)], check=True)


def generate(cn, client_dir: str | Path, take: str, work: str | Path, out: str | Path, pdf: bool = True) -> dict[str, Any]:
    client_dir = Path(client_dir)
    work = Path(work)
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    client = read_yaml(client_dir / "take.yaml")["client"]
    data = collect(cn, client, take, client_dir, work)
    checks = verify_numbers(data)
    html_text = html_report(data, checks)
    html_path = out / f"{client}-{take}-data-integrity-report.html"
    docx_path = out / f"{client}-{take}-data-integrity-report.docx"
    pdf_path = out / f"{client}-{take}-data-integrity-report.pdf"
    html_path.write_text(html_text, encoding="utf-8", newline="\r\n")
    write_docx(data, checks, docx_path)
    pdf_note = None
    if pdf:                     # needs a local Edge (laptop); Fabric notebooks keep HTML and DOCX
        try:
            render_pdf(html_path, pdf_path)
        except Exception as e:  # noqa: BLE001 - the report is complete without the PDF copy
            pdf_note = f"PDF not rendered: {type(e).__name__}"
    return {"html": str(html_path), "docx": str(docx_path), "pdf": str(pdf_path) if pdf and not pdf_note else None, "pdf_note": pdf_note,
            "checks": checks, "answer_key": answer_summary_text(data), "modules": len(data["modules"]), "issues": len(data["issues"])}
